import json
import sys
import tempfile
import types
import unittest
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from openpyxl import load_workbook

import takeout_organizer
from takeout_organizer import (
    _album_name,
    _media_datetime,
    _metadata_datetime,
    _safe_member_path,
    discover_archives,
    extract_archives,
    organize_extracted,
    organize_takeout,
)
from date_helpers import (
    find_thm_sidecar,
    get_exif_datetime,
    get_video_created_datetime,
    parse_datetime_from_filename,
)


class FakeVariable:
    def __init__(self):
        self.value = ""

    def get(self):
        return self.value

    def set(self, value):
        self.value = value


class FakeWidget:
    def __init__(self, *args, **kwargs):
        self.options = kwargs
        self.content = ""

    def pack(self, **kwargs):
        pass

    def grid(self, **kwargs):
        pass

    def columnconfigure(self, *args, **kwargs):
        pass

    def rowconfigure(self, *args, **kwargs):
        pass

    def configure(self, **kwargs):
        self.options.update(kwargs)

    def start(self, *args):
        pass

    def stop(self):
        pass

    def insert(self, index, value):
        self.content += value

    def delete(self, start, end):
        self.content = ""

    def see(self, index):
        pass

    def yview(self, *args):
        pass

    def set(self, *args):
        pass

    def tag_configure(self, *args, **kwargs):
        pass


class FakeRoot:
    def __init__(self):
        self.callbacks = []
        self.mainloop_called = False

    def title(self, value):
        pass

    def minsize(self, *args):
        pass

    def after(self, delay, callback):
        self.callbacks.append((delay, callback))

    def mainloop(self):
        self.mainloop_called = True


class TakeoutOrganizerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def _write_media(self, relative_path, content=b"media", metadata=None):
        media = self.root / relative_path
        media.parent.mkdir(parents=True, exist_ok=True)
        media.write_bytes(content)
        if metadata is not None:
            media.with_name(media.name + ".json").write_text(json.dumps(metadata), encoding="utf-8")
        return media

    def _new_gui(self):
        fake_tk = types.SimpleNamespace(StringVar=FakeVariable)
        fake_ttk = types.SimpleNamespace(
            Frame=FakeWidget,
            Label=FakeWidget,
            Entry=FakeWidget,
            Button=FakeWidget,
            Progressbar=FakeWidget,
            Scrollbar=FakeWidget,
        )
        fake_tk.Text = FakeWidget
        with patch.dict(sys.modules, {"tkinter": fake_tk}), patch.object(takeout_organizer, "ttk", fake_ttk):
            root = FakeRoot()
            window = takeout_organizer.OrganizerWindow(root)
        return root, window

    def test_filename_date_precedes_takeout_timestamps(self):
        timestamp = str(int(datetime(2026, 10, 3, tzinfo=timezone.utc).timestamp()))
        media = self._write_media(
            "input/Takeout/Google Photos/2026 - trip/IMG-20200101-WA0001.jpg",
            metadata={"photoTakenTime": {"timestamp": timestamp}},
        )
        self.assertEqual(_metadata_datetime(media), datetime(2026, 10, 3, tzinfo=timezone.utc).replace(tzinfo=None))
        self.assertEqual(_media_datetime(media), datetime(2020, 1, 1))

    def test_creation_time_used_when_photo_taken_time_is_invalid(self):
        media = self._write_media(
            "creation/photo.jpg",
            metadata={"photoTakenTime": {"timestamp": "bad"}, "creationTime": {"timestamp": "1577836800"}},
        )
        self.assertEqual(_metadata_datetime(media, "creationTime"), datetime(2020, 1, 1))

    def test_google_metadata_year_outside_allowed_range_is_ignored(self):
        media = self._write_media(
            "creation/too-old.jpg",
            metadata={"creationTime": {"timestamp": "-5364662400"}},
        )
        self.assertIsNone(_metadata_datetime(media, "creationTime"))

    def test_non_object_and_missing_timestamp_sidecars_are_ignored(self):
        media = self._write_media("metadata/invalid.jpg")
        sidecar = media.with_name("invalid.jpg.json")
        sidecar.write_text("[]", encoding="utf-8")
        self.assertIsNone(_metadata_datetime(media))
        sidecar.write_text(json.dumps({"photoTakenTime": []}), encoding="utf-8")
        self.assertIsNone(_metadata_datetime(media))

    def test_filename_capture_date_precedes_takeout_creation_time(self):
        media = self._write_media(
            "creation/IMG-20211011-WA0008.jpg",
            metadata={"creationTime": {"timestamp": "1577836800"}},
        )
        self.assertEqual(_media_datetime(media).date(), datetime(2021, 10, 11).date())

    def test_takeout_creation_time_is_used_before_filesystem_mtime(self):
        media = self._write_media(
            "creation/unknown.jpg",
            metadata={"creationTime": {"timestamp": "1577836800"}},
        )
        self.assertEqual(_media_datetime(media), datetime(2020, 1, 1))

    def test_malformed_sidecar_is_ignored_and_supplemental_sidecar_is_read(self):
        media = self._write_media("metadata/photo.jpg")
        media.with_name("photo.jpg.json").write_text("{broken", encoding="utf-8")
        media.with_name("photo.jpg.supplemental-metadata.json").write_text(
            json.dumps({"photoTakenTime": {"timestamp": "1577836800"}}), encoding="utf-8"
        )
        self.assertEqual(_metadata_datetime(media), datetime(2020, 1, 1))

    def test_out_of_range_sidecar_timestamp_falls_through_to_supplemental_file(self):
        media = self._write_media("metadata/overflow.jpg")
        media.with_name("overflow.jpg.json").write_text(
            json.dumps({"photoTakenTime": {"timestamp": "not-an-epoch"}}), encoding="utf-8"
        )
        media.with_name("overflow.jpg.supplemental-metadata.json").write_text(
            json.dumps({"photoTakenTime": {"timestamp": "1577836800"}}), encoding="utf-8"
        )
        self.assertEqual(_metadata_datetime(media), datetime(2020, 1, 1))

    def test_existing_filename_timestamp_formats_remain_supported(self):
        cases = (
            ("vlcsnap-2011-05-13-15h20m43s172", datetime(2011, 5, 13, 15, 20, 43)),
            ("WIN_20200330_09_31_37_Pro", datetime(2020, 3, 30, 9, 31, 37)),
            ("20230315_182317", datetime(2023, 3, 15, 18, 23, 17)),
            ("20230315182317", datetime(2023, 3, 15, 18, 23, 17)),
            ("2023-03-15_18-23-17", datetime(2023, 3, 15, 18, 23, 17)),
            ("WP_20140407_005", datetime(2014, 4, 7)),
            ("IMG-20211011-WA0008", datetime(2021, 10, 11)),
        )
        for filename, expected in cases:
            with self.subTest(filename=filename):
                self.assertEqual(parse_datetime_from_filename(filename), expected)

    def test_user_filename_examples_and_year_bounds(self):
        cases = (
            ("2007-01-01 00.13.46 2007 01 01", datetime(2007, 1, 1, 0, 13, 46)),
            ("27112009 27 11(nov) 2009", datetime(2009, 11, 27)),
            ("01012010024", datetime(2010, 1, 1)),
            ("WhatsApp Image 2021-04-29 at 18.46.14 (3)", datetime(2021, 4, 29, 18, 46, 14)),
            ("Video2007730423257", datetime(2007, 7, 30)),
            ("IMG_19500101", datetime(1950, 1, 1)),
            ("IMG_21001231", datetime(2100, 12, 31)),
        )
        for filename, expected in cases:
            with self.subTest(filename=filename):
                self.assertEqual(parse_datetime_from_filename(filename), expected)
        self.assertIsNone(parse_datetime_from_filename("IMG_10290101"))
        self.assertIsNone(parse_datetime_from_filename("IMG_06590101"))

    def test_invalid_filename_dates_and_media_without_exif_return_none(self):
        self.assertIsNone(parse_datetime_from_filename("vlcsnap-2023-99-99-99h99m99s"))
        self.assertIsNone(parse_datetime_from_filename("WP_20231340"))
        media = self.root / "no-exif.jpg"
        Image.new("RGB", (1, 1)).save(media)
        self.assertIsNone(get_exif_datetime(media))

    def test_filename_date_precedes_exif(self):
        media = self.root / "exif" / "IMG-20200101.jpg"
        media.parent.mkdir()
        image = Image.new("RGB", (1, 1))
        exif = image.getexif()
        exif[36867] = "2019:12:31 23:59:58"
        image.save(media, exif=exif)
        self.assertEqual(_media_datetime(media), datetime(2020, 1, 1))

    def test_exif_date_precedes_google_photos_creation_time(self):
        media = self.root / "exif" / "unknown.jpg"
        media.parent.mkdir()
        image = Image.new("RGB", (1, 1))
        exif = image.getexif()
        exif[36867] = "2019:12:31 23:59:58"
        image.save(media, exif=exif)
        media.with_name("unknown.jpg.json").write_text(
            json.dumps({"creationTime": {"timestamp": "1577836800"}}), encoding="utf-8"
        )
        self.assertEqual(_media_datetime(media), datetime(2019, 12, 31, 23, 59, 58))

    def test_video_thm_exif_precedes_google_photos_creation_time(self):
        video = self._write_media(
            "thm/clip.mp4",
            metadata={"creationTime": {"timestamp": "1577836800"}},
        )
        thm = video.with_suffix(".THM")
        image = Image.new("RGB", (1, 1))
        exif = image.getexif()
        exif[36867] = "2018:07:06 05:04:03"
        image.save(thm, format="JPEG", exif=exif)
        self.assertEqual(find_thm_sidecar(video), thm)
        self.assertEqual(get_exif_datetime(thm), datetime(2018, 7, 6, 5, 4, 3))
        self.assertEqual(_media_datetime(video), datetime(2018, 7, 6, 5, 4, 3))
        result = organize_extracted([self.root / "thm"], self.root / "output")
        self.assertEqual(result.report_rows[0]["Found date from .THM"], datetime(2018, 7, 6, 5, 4, 3))

    def test_video_media_created_metadata_precedes_google_json(self):
        video = self._write_media(
            "video/DSCN1943.MOV",
            metadata={"creationTime": {"timestamp": "1577836800"}},
        )
        media_created = datetime(2007, 7, 30, 12, 34, 56)
        ffprobe_result = types.SimpleNamespace(
            stdout=json.dumps({"format": {"tags": {"creation_time": "2007-07-30T12:34:56"}}})
        )
        with patch("date_helpers.shutil.which", return_value="ffprobe"), patch(
            "date_helpers.subprocess.run", return_value=ffprobe_result
        ) as run:
            self.assertEqual(get_video_created_datetime(video), media_created)
            run.assert_called_once()
        with patch("takeout_organizer.get_video_created_datetime", return_value=media_created):
            self.assertEqual(_media_datetime(video), media_created)
            result = organize_extracted([video.parent], self.root / "organized")
        self.assertEqual(result.report_rows[0]["Found date in media created"], media_created)
        self.assertTrue((self.root / "organized/2007/2007 July/DSCN1943.MOV").is_file())

    def test_ffprobe_absence_falls_back_without_error(self):
        video = self._write_media("video/no-tag.mov")
        with patch("date_helpers.shutil.which", return_value=None):
            self.assertIsNone(get_video_created_datetime(video))

    def test_ffprobe_reads_stream_date_and_ignores_bad_or_out_of_range_tags(self):
        video = self._write_media("video/stream-date.mov")
        probe_output = {
            "format": {"tags": {"creation_time": "not-a-date"}},
            "streams": [{"tags": {"creation_time": "2007:07:30 12:34:56"}}],
        }
        with patch("date_helpers.shutil.which", return_value="ffprobe"), patch(
            "date_helpers.subprocess.run",
            return_value=types.SimpleNamespace(stdout=json.dumps(probe_output)),
        ):
            self.assertEqual(get_video_created_datetime(video), datetime(2007, 7, 30, 12, 34, 56))

        probe_output = {"format": {"tags": {"creation_time": "1800-01-01T00:00:00"}}}
        with patch("date_helpers.shutil.which", return_value="ffprobe"), patch(
            "date_helpers.subprocess.run",
            return_value=types.SimpleNamespace(stdout=json.dumps(probe_output)),
        ):
            self.assertIsNone(get_video_created_datetime(video))

        with patch("date_helpers.shutil.which", return_value="ffprobe"), patch(
            "date_helpers.subprocess.run", side_effect=TimeoutError("probe timed out")
        ):
            self.assertIsNone(get_video_created_datetime(video))

    def test_collage_media_is_copied_to_year_collage_folder(self):
        extraction = self.root / "extract"
        source = extraction / "Takeout/Google Photos/Family/COLLAGE photos.jpg"
        source.parent.mkdir(parents=True)
        source.write_bytes(b"collage")
        source.with_name(source.name + ".json").write_text(
            json.dumps({"creationTime": {"timestamp": "1577836800"}}), encoding="utf-8"
        )
        result = organize_extracted([extraction], self.root / "output")
        self.assertTrue((self.root / "output/2020/collage/COLLAGE photos.jpg").is_file())
        self.assertEqual(result.album_copies, 1)

    def test_filename_parser_and_no_date(self):
        dated = self._write_media("plain/IMG-20211011-WA0008.jpg")
        dated.touch()
        self.assertEqual(_media_datetime(dated).date(), datetime(2021, 10, 11).date())
        undated = self._write_media("plain/unknown.jpg")
        self.assertIsNone(_media_datetime(undated))

    def test_album_detection_skips_date_directories(self):
        root = self.root / "extract"
        path = root / "Takeout" / "Google Photos" / "2026 - trip" / "2026 October" / "photo.jpg"
        self.assertEqual(_album_name(path, root), "2026 - trip")
        month_path = root / "Takeout" / "Google Photos" / "2026" / "2026 October" / "photo.jpg"
        self.assertIsNone(_album_name(month_path, root))
        photos_from_year = root / "Takeout" / "Google Photos" / "Photos from 2026" / "photo.jpg"
        self.assertIsNone(_album_name(photos_from_year, root))
        generic_album = root / "Album name" / "photo.jpg"
        self.assertEqual(_album_name(generic_album, root), "Album name")

    def test_safe_member_path_rejects_traversal(self):
        with self.assertRaises(ValueError):
            _safe_member_path(self.root / "extract", "../outside.txt")
        with self.assertRaises(ValueError):
            _safe_member_path(self.root / "extract", "C:/outside.txt")
        with self.assertRaises(ValueError):
            _safe_member_path(self.root / "extract", ".")
        self.assertEqual(
            _safe_member_path(self.root / "extract", "Takeout/photo.jpg").name,
            "photo.jpg",
        )
        sanitized = _safe_member_path(self.root / "extract", "Takeout/2022 - Album:/photo.jpg")
        self.assertIn("2022 - Album_", str(sanitized))

    def test_archive_with_windows_invalid_album_name_is_extracted_safely(self):
        archive_path = self.root / "invalid-folder-name.zip"
        with zipfile.ZipFile(archive_path, "w") as archive:
            archive.writestr("Takeout/Google Photos/2022 - Album:/photo.jpg", b"photo")
        messages = []
        [extraction_root] = extract_archives([archive_path], self.root / "tmp", messages.append)
        extracted_file = extraction_root / "Takeout/Google Photos/2022 - Album_/photo.jpg"
        self.assertEqual(extracted_file.read_bytes(), b"photo")
        self.assertTrue(any("Sanitized Windows-invalid path" in message for message in messages))

    def test_safe_member_path_rejects_resolved_escape(self):
        destination = self.root / "extract"
        escaped_target = self.root / "outside.txt"
        original_resolve = Path.resolve

        def resolve_with_escape(path, *args, **kwargs):
            if path == destination / "inside.txt":
                return escaped_target
            return original_resolve(path, *args, **kwargs)

        with patch.object(Path, "resolve", autospec=True, side_effect=resolve_with_escape):
            with self.assertRaisesRegex(ValueError, "Unsafe path"):
                _safe_member_path(destination, "inside.txt")

    def test_archive_discovery_is_recursive_and_sorted(self):
        (self.root / "b.zip").touch()
        nested = self.root / "nested"
        nested.mkdir()
        (nested / "a.ZIP").touch()
        self.assertEqual([path.name for path in discover_archives(self.root)], ["a.ZIP", "b.zip"])

    def test_extract_archives_keeps_sources_and_separates_same_paths(self):
        archive_paths = []
        for name, contents in (("part1.zip", b"one"), ("part2.zip", b"two")):
            archive_path = self.root / name
            with zipfile.ZipFile(archive_path, "w") as archive:
                archive.writestr("Takeout/Google Photos/", "")
                archive.writestr("Takeout/Google Photos/photo.jpg", contents)
            archive_paths.append(archive_path)
        roots = extract_archives(archive_paths, self.root / "tmp")
        self.assertEqual(len(roots), 2)
        self.assertEqual((roots[0] / "Takeout/Google Photos/photo.jpg").read_bytes(), b"one")
        self.assertEqual((roots[1] / "Takeout/Google Photos/photo.jpg").read_bytes(), b"two")
        self.assertTrue(all(path.exists() for path in archive_paths))

    def test_invalid_zip_and_empty_archive_list_report_errors(self):
        with self.assertRaises(ValueError):
            extract_archives([], self.root / "tmp")
        invalid_zip = self.root / "bad.zip"
        invalid_zip.write_text("not a zip", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Not a valid ZIP"):
            extract_archives([invalid_zip], self.root / "tmp")

    def test_fatal_extraction_failure_cleans_partial_run_folder(self):
        archive_directory = self.root / "archives"
        archive_directory.mkdir()
        (archive_directory / "invalid.zip").write_text("not a ZIP", encoding="utf-8")
        temporary = self.root / "tmp"
        with self.assertRaisesRegex(ValueError, "Not a valid ZIP"):
            organize_takeout(archive_directory, temporary, self.root / "output")
        self.assertEqual(list(temporary.iterdir()), [])

    def test_bad_crc_member_is_logged_and_later_media_is_processed(self):
        archive_path = self.root / "partially-corrupt.zip"
        with zipfile.ZipFile(archive_path, "w") as archive:
            archive.writestr("Takeout/Google Photos/Album/broken.jpg", b"CORRUPT", compress_type=zipfile.ZIP_STORED)
            archive.writestr("Takeout/Google Photos/Album/IMG-20200101-good.jpg", b"good", compress_type=zipfile.ZIP_STORED)
        archive_bytes = archive_path.read_bytes()
        archive_path.write_bytes(archive_bytes.replace(b"CORRUPT", b"XORRUPT"))
        messages = []
        result = organize_takeout(self.root, self.root / "tmp", self.root / "output", messages.append)
        self.assertEqual(len(result.errors), 1)
        self.assertIn("broken.jpg", result.errors[0])
        self.assertIn("[ERROR]", result.errors[0])
        self.assertTrue(any("Processing media" in message for message in messages))
        self.assertTrue((self.root / "output/2020/2020 January/IMG-20200101-good.jpg").is_file())

    def test_end_to_end_takeout_layout_and_collision_handling(self):
        archive_path = self.root / "takeout.zip"
        timestamp = str(int(datetime(2026, 10, 3, tzinfo=timezone.utc).timestamp()))
        with zipfile.ZipFile(archive_path, "w") as archive:
            archive.writestr("Takeout/Google Photos/2026 - trip/photo.jpg", b"first")
            archive.writestr(
                "Takeout/Google Photos/2026 - trip/photo.jpg.json",
                json.dumps({"creationTime": {"timestamp": timestamp}}),
            )
            archive.writestr("Takeout/Google Photos/2026 - trip/clip.mp4", b"video")
            archive.writestr(
                "Takeout/Google Photos/2026 - trip/clip.mp4.json",
                json.dumps({"creationTime": {"timestamp": timestamp}}),
            )
            archive.writestr("Takeout/Google Photos/2026 - trip/metadata.json", "{}")
            archive.writestr("Takeout/Google Photos/2026 - trip/no-date.jpg", b"unknown")
        output = self.root / "output"
        temporary = self.root / "temporary"
        temporary.mkdir()
        unrelated_temp_file = temporary / "keep-this-file.txt"
        unrelated_temp_file.write_text("not created by GoogLi", encoding="utf-8")
        existing_month_file = output / "2026" / "2026 October" / "photo.jpg"
        existing_month_file.parent.mkdir(parents=True)
        existing_month_file.write_bytes(b"other")
        messages = []
        result = organize_takeout(self.root, temporary, output, messages.append)
        month = output / "2026" / "2026 October"
        album = output / "2026" / "2026 - trip"
        self.assertEqual(result.archives, 1)
        self.assertEqual(result.media_found, 3)
        self.assertEqual(result.month_copies, 1)
        self.assertEqual(result.album_copies, 2)
        self.assertEqual(result.unknown_date_copies, 1)
        self.assertTrue((month / "photo.jpg").is_file())
        self.assertTrue((month / "clip.mp4").is_file())
        self.assertTrue((album / "photo.jpg").is_file())
        self.assertTrue((output / "unknown date" / "no-date.jpg").is_file())
        self.assertFalse((month / "metadata.json").exists())
        self.assertTrue(archive_path.exists())
        self.assertEqual(existing_month_file.read_bytes(), b"other")
        self.assertEqual(list(temporary.iterdir()), [unrelated_temp_file])
        self.assertTrue(any(message.startswith("Cleaning temporary extraction data:") for message in messages))
        self.assertTrue(result.report_path.is_file())
        workbook = load_workbook(result.report_path, read_only=True)
        try:
            worksheet = workbook["Media audit"]
            headers = [cell.value for cell in worksheet[1]]
            self.assertEqual(headers, list(takeout_organizer.REPORT_COLUMNS))
            rows = {Path(row[1]).name: row for row in worksheet.iter_rows(min_row=2, values_only=True)}
            photo_row = rows["photo.jpg"]
            self.assertEqual(photo_row[6], datetime(2026, 10, 3))
            self.assertIn("takeout.zip :: Takeout/Google Photos/2026 - trip/photo.jpg", photo_row[7])
            self.assertIn(str(existing_month_file), photo_row[8])
            self.assertTrue(photo_row[9].startswith("Yes (month; same filename and size)"))
            self.assertEqual(rows["no-date.jpg"][9], "No")
        finally:
            workbook.close()

    def test_same_named_files_get_unique_destinations(self):
        extraction = self.root / "extract"
        first = extraction / "Takeout/Google Photos/Album A/same.jpg"
        second = extraction / "Takeout/Google Photos/Album B/same.jpg"
        first.parent.mkdir(parents=True)
        second.parent.mkdir(parents=True)
        first.write_bytes(b"first")
        second.write_bytes(b"second")
        first.with_name("same.jpg.json").write_text(
            json.dumps({"photoTakenTime": {"timestamp": "1577836800"}}), encoding="utf-8"
        )
        second.with_name("same.jpg.json").write_text(
            json.dumps({"photoTakenTime": {"timestamp": "1577836800"}}), encoding="utf-8"
        )
        result = organize_extracted([extraction], self.root / "output")
        month = self.root / "output/2020/2020 January"
        self.assertEqual(result.month_copies, 2)
        self.assertEqual({path.name for path in month.iterdir()}, {"same.jpg", "same (2).jpg"})

    def test_same_filename_and_size_is_not_copied_twice(self):
        extractions = [self.root / "extract1", self.root / "extract2"]
        for extraction in extractions:
            media = extraction / "Takeout/Google Photos/Album/same.jpg"
            media.parent.mkdir(parents=True, exist_ok=True)
            media.write_bytes(b"same bytes")
            media.with_name("same.jpg.json").write_text(
                json.dumps({"creationTime": {"timestamp": "1577836800"}}), encoding="utf-8"
            )
        result = organize_extracted(extractions, self.root / "output")
        month = self.root / "output/2020/2020 January"
        self.assertEqual(result.month_copies, 1)
        self.assertEqual(result.album_copies, 1)
        self.assertEqual(len(list(month.iterdir())), 1)
        self.assertEqual(result.report_rows[0]["If the file already existed"], "No")
        self.assertEqual(
            result.report_rows[1]["If the file already existed"],
            "Yes (month, album; same filename and size)",
        )

    def test_unreadable_media_is_copied_to_unknown_date(self):
        extraction = self.root / "extract"
        media = extraction / "Takeout/Google Photos/photo.jpg"
        media.parent.mkdir(parents=True)
        media.write_bytes(b"media")
        with patch("takeout_organizer._media_datetime", side_effect=ValueError("bad date")):
            result = organize_extracted([extraction], self.root / "output")
        self.assertEqual(result.media_found, 1)
        self.assertEqual(result.unknown_date_copies, 1)
        self.assertEqual(result.month_copies, 0)
        self.assertTrue((self.root / "output/unknown date/photo.jpg").is_file())
        with patch("takeout_organizer._media_datetime", side_effect=ValueError("bad date")):
            duplicate_result = organize_extracted([extraction], self.root / "output")
        self.assertEqual(duplicate_result.unknown_date_copies, 0)
        self.assertEqual(
            duplicate_result.report_rows[0]["If the file already existed"],
            "Yes (unknown date; same filename and size)",
        )

    def test_gui_rejects_missing_and_nonexistent_input_paths(self):
        _, window = self._new_gui()
        with patch("takeout_organizer.messagebox.showerror") as showerror:
            window._start()
            showerror.assert_called_once_with(
                "Missing folder", "Select the ZIP, temporary, and output folders first."
            )
            for variable in window.path_vars.values():
                variable.set(str(self.root / "not-created"))
            window._start()
            self.assertEqual(showerror.call_count, 2)

    def test_gui_browse_sets_selected_directory(self):
        _, window = self._new_gui()
        variable = next(iter(window.path_vars.values()))
        with patch("takeout_organizer.filedialog.askdirectory", return_value=str(self.root)):
            window._browse(variable)
        self.assertEqual(variable.get(), str(self.root))

    def test_gui_worker_success_and_error_events(self):
        class InlineThread:
            def __init__(self, target, args, daemon):
                self.target = target
                self.args = args

            def start(self):
                self.target(*self.args)

        archive_directory = self.root / "archives"
        archive_directory.mkdir()
        root, window = self._new_gui()
        for label, variable in window.path_vars.items():
            variable.set(str(archive_directory if label == "ZIP folder" else self.root / label))
        result = takeout_organizer.SortResult(archives=1, media_found=2, month_copies=2)
        with patch("takeout_organizer.threading.Thread", InlineThread), \
                patch("takeout_organizer.organize_takeout", return_value=result), \
                patch("takeout_organizer.messagebox.showinfo") as showinfo:
            window._start()
            window._poll_events()
        showinfo.assert_called_once()
        self.assertEqual(window.status.options["text"], "Organization complete.")
        self.assertEqual(window.start_button.options["state"], "normal")

        _, failed_window = self._new_gui()
        for label, variable in failed_window.path_vars.items():
            variable.set(str(archive_directory if label == "ZIP folder" else self.root / label))
        with patch("takeout_organizer.threading.Thread", InlineThread), \
                patch("takeout_organizer.organize_takeout", side_effect=RuntimeError("broken ZIP")), \
                patch("takeout_organizer.messagebox.showerror") as showerror:
            failed_window._start()
            failed_window._poll_events()
        showerror.assert_called_once_with("Organization failed", "broken ZIP")

    def test_gui_poll_waits_when_no_worker_event_is_ready(self):
        root, window = self._new_gui()
        window._poll_events()
        self.assertEqual(root.callbacks[-1][0], 150)

    def test_gui_activity_log_displays_progress_events(self):
        root, window = self._new_gui()
        window.events.put(("log", "Extracting ZIP 1/1: export.zip"))
        window.events.put(("log", "Processing media 1: export.zip :: photo.jpg"))
        window._poll_events()
        self.assertIn("Extracting ZIP 1/1: export.zip", window.activity_log.content)
        self.assertIn("Processing media 1: export.zip :: photo.jpg", window.activity_log.content)
        self.assertEqual(window.status.options["text"], "Processing media 1: export.zip :: photo.jpg")
        self.assertEqual(root.callbacks[-1][0], 50)

    def test_launch_gui_starts_tk_event_loop(self):
        root = FakeRoot()
        with patch("takeout_organizer.Tk", return_value=root), patch(
            "takeout_organizer.OrganizerWindow"
        ):
            takeout_organizer.launch_gui()
        self.assertTrue(root.mainloop_called)


if __name__ == "__main__":
    unittest.main()