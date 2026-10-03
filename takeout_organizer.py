"""Organize Google Photos Takeout archives into year, month, and album folders."""

from __future__ import annotations

import json
import os
import queue
import re
import shutil
import tempfile
import threading
import zipfile
import zlib
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from tkinter import Tk, filedialog, messagebox, ttk

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from date_helpers import (
    find_thm_sidecar,
    get_exif_datetime,
    get_video_created_datetime,
    is_valid_media_datetime,
    parse_datetime_from_filename,
)


IMAGE_EXTENSIONS = {
    ".jpg", ".jpeg", ".png", ".gif", ".tif", ".tiff", ".webp", ".heic",
    ".heif", ".bmp", ".dng", ".cr2", ".nef", ".arw", ".orf", ".rw2",
}
VIDEO_EXTENSIONS = {
    ".mp4", ".mov", ".m4v", ".avi", ".mkv", ".3gp", ".3g2", ".wmv",
    ".mpg", ".mpeg", ".mts", ".m2ts", ".mod", ".webm",
}
MEDIA_EXTENSIONS = IMAGE_EXTENSIONS | VIDEO_EXTENSIONS
MONTH_NAMES = (
    "January", "February", "March", "April", "May", "June", "July",
    "August", "September", "October", "November", "December",
)
DATE_FOLDER = re.compile(
    r"^(?:\d{4}(?:[-_]\d{2})?|\d{4}\s+(?:January|February|March|April|May|June|July|August|September|October|November|December)(?:-\d+-\d+)?|Photos from \d{4})$",
    re.IGNORECASE,
)


@dataclass
class SortResult:
    archives: int = 0
    media_found: int = 0
    dated_media: int = 0
    month_copies: int = 0
    album_copies: int = 0
    unknown_date_copies: int = 0
    report_rows: list[dict[str, object]] = field(default_factory=list)
    report_path: Path | None = None
    errors: list[str] = field(default_factory=list)


def discover_archives(archive_directory: Path) -> list[Path]:
    """Find ZIP archives recursively, in stable order."""
    return sorted(
        (path for path in archive_directory.rglob("*")
         if path.is_file() and path.suffix.lower() == ".zip"),
        key=lambda path: (path.name.casefold(), str(path).casefold()),
    )


def _sanitize_windows_component(component: str) -> str:
    sanitized = re.sub(r'[<>:"|?*]', "_", component).rstrip(" .")
    if not sanitized:
        sanitized = "_"
    device_name = sanitized.split(".", 1)[0].upper()
    if device_name in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}:
        sanitized = f"_{sanitized}"
    return sanitized


def _safe_member_path(destination: Path, member_name: str) -> Path:
    normalized = member_name.replace("\\", "/")
    member = PurePosixPath(normalized)
    if (not member.parts or member.is_absolute()
            or any(part in ("..", "") for part in member.parts)
            or re.match(r"^[A-Za-z]:", member.parts[0])):
        raise ValueError(f"Unsafe path in ZIP archive: {member_name}")
    safe_parts = tuple(_sanitize_windows_component(part) for part in member.parts)
    target = destination.joinpath(*safe_parts).resolve()
    if os.path.commonpath((str(destination.resolve()), str(target))) != str(destination.resolve()):
        raise ValueError(f"Unsafe path in ZIP archive: {member_name}")
    return target


def _extract_archive(
    archive_path: Path,
    destination: Path,
    progress_callback: Callable[[str], None] | None = None,
) -> None:
    if not zipfile.is_zipfile(archive_path):
        raise ValueError(f"Not a valid ZIP archive: {archive_path}")

    destination.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(archive_path) as archive:
        for member in archive.infolist():
            target = None
            try:
                target = _safe_member_path(destination, member.filename)
                normalized_member = member.filename.replace("\\", "/")
                member_parts = PurePosixPath(normalized_member).parts
                safe_parts = tuple(_sanitize_windows_component(part) for part in member_parts)
                safe_member = "/".join(safe_parts)
                if progress_callback:
                    if safe_parts != member_parts:
                        progress_callback(f"Sanitized Windows-invalid path: {member.filename} -> {safe_member}")
                    progress_callback(f"Extracting: {archive_path.name} :: {member.filename}")
                if member.is_dir():
                    target.mkdir(parents=True, exist_ok=True)
                    continue
                target.parent.mkdir(parents=True, exist_ok=True)
                with archive.open(member) as source, target.open("wb") as output:
                    shutil.copyfileobj(source, output)
            except (OSError, zipfile.BadZipFile, zlib.error) as error:
                if target is not None:
                    try:
                        target.unlink(missing_ok=True)
                    except OSError:
                        pass
                message = f"[ERROR] Failed extracting ZIP '{archive_path}', member '{member.filename}': {error}; skipped"
                if progress_callback:
                    progress_callback(message)


def extract_archives(
    archive_paths: list[Path],
    temp_directory: Path,
    progress_callback: Callable[[str], None] | None = None,
) -> list[Path]:
    """Extract archives into a unique run folder without modifying the ZIPs."""
    if not archive_paths:
        raise ValueError("No .zip files were found in the selected directory.")
    temp_directory.mkdir(parents=True, exist_ok=True)
    run_directory = Path(tempfile.mkdtemp(prefix="takeout_import_", dir=temp_directory))
    extracted_roots = []
    try:
        for index, archive_path in enumerate(archive_paths, start=1):
            archive_root = run_directory / f"{index:04d}"
            if progress_callback:
                progress_callback(f"Extracting ZIP {index}/{len(archive_paths)}: {archive_path}")
            _extract_archive(archive_path, archive_root, progress_callback)
            extracted_roots.append(archive_root)
    except Exception:
        shutil.rmtree(run_directory, ignore_errors=True)
        raise
    return extracted_roots


def _sidecar_paths(media_path: Path) -> tuple[Path, ...]:
    return (
        media_path.with_name(media_path.name + ".json"),
        media_path.with_name(media_path.name + ".supplemental-metadata.json"),
        media_path.with_suffix(".json"),
    )


def _metadata_datetime(media_path: Path, field: str = "photoTakenTime") -> datetime | None:
    for sidecar in _sidecar_paths(media_path):
        if not sidecar.is_file():
            continue
        try:
            metadata = json.loads(sidecar.read_text(encoding="utf-8-sig"))
        except (OSError, UnicodeError, json.JSONDecodeError):
            continue
        if not isinstance(metadata, dict):
            continue
        date_data = metadata.get(field)
        timestamp = date_data.get("timestamp") if isinstance(date_data, dict) else None
        if timestamp is None:
            continue
        try:
            parsed = datetime.fromtimestamp(float(timestamp), tz=timezone.utc).replace(tzinfo=None)
            if is_valid_media_datetime(parsed):
                return parsed
        except (ValueError, TypeError, OverflowError, OSError):
            continue
    return None


def _media_datetime(media_path: Path) -> datetime | None:
    filename_date = parse_datetime_from_filename(media_path.stem)
    exif_date = get_exif_datetime(media_path)
    thm_date = None
    if media_path.suffix.lower() in VIDEO_EXTENSIONS:
        thm_path = find_thm_sidecar(media_path)
        if thm_path:
            thm_date = get_exif_datetime(thm_path)
    media_created_date = None
    if not (filename_date or exif_date or thm_date) and media_path.suffix.lower() in VIDEO_EXTENSIONS:
        media_created_date = get_video_created_datetime(media_path)
    google_date = _metadata_datetime(media_path, "creationTime") or _metadata_datetime(media_path)
    return filename_date or exif_date or thm_date or media_created_date or google_date


def _album_name(media_path: Path, extraction_root: Path) -> str | None:
    relative_parts = media_path.relative_to(extraction_root).parts[:-1]
    google_photos_index = next(
        (index for index, part in enumerate(relative_parts)
         if part.casefold() == "google photos"),
        None,
    )
    if google_photos_index is not None:
        candidate_parts = relative_parts[google_photos_index + 1:]
    else:
        candidate_parts = tuple(
            part for part in relative_parts
            if part.casefold() != "takeout" and not part.casefold().startswith("takeout_")
        )
    return next(
        (part for part in candidate_parts if not DATE_FOLDER.fullmatch(part)),
        None,
    )


def _is_collage(media_path: Path, extraction_root: Path) -> bool:
    return any("collage" in part.casefold() for part in media_path.relative_to(extraction_root).parts)


def _copy_without_overwrite(source: Path, destination: Path) -> tuple[Path, bool]:
    destination.parent.mkdir(parents=True, exist_ok=True)
    candidate = destination
    counter = 2
    while candidate.exists():
        if candidate.is_file() and candidate.stat().st_size == source.stat().st_size:
            return candidate, True
        candidate = destination.with_name(f"{destination.stem} ({counter}){destination.suffix}")
        counter += 1
    shutil.copy2(source, candidate)
    return candidate, False


def organize_extracted(
    extracted_roots: list[Path],
    output_directory: Path,
    source_archives: dict[Path, Path] | None = None,
    progress_callback: Callable[[str], None] | None = None,
) -> SortResult:
    """Copy recognized media into year/month folders and year-level albums."""
    output_directory.mkdir(parents=True, exist_ok=True)
    result = SortResult()
    for extraction_root in extracted_roots:
        archive_path = (source_archives or {}).get(extraction_root)
        for media_path in sorted(extraction_root.rglob("*"), key=lambda path: str(path).casefold()):
            if not media_path.is_file() or media_path.suffix.lower() not in MEDIA_EXTENSIONS:
                continue
            result.media_found += 1
            if progress_callback:
                location = f"{archive_path.name} :: " if archive_path else ""
                progress_callback(f"Processing media {result.media_found}: {location}{media_path.relative_to(extraction_root)}")
            filename_date = parse_datetime_from_filename(media_path.stem)
            exif_date = get_exif_datetime(media_path)
            thm_date = None
            if media_path.suffix.lower() in VIDEO_EXTENSIONS:
                thm_path = find_thm_sidecar(media_path)
                if thm_path:
                    thm_date = get_exif_datetime(thm_path)
            media_created_date = None
            if not (filename_date or exif_date or thm_date) and media_path.suffix.lower() in VIDEO_EXTENSIONS:
                media_created_date = get_video_created_datetime(media_path)
            creation_date = (
                _metadata_datetime(media_path, "creationTime")
                or _metadata_datetime(media_path, "photoTakenTime")
            )
            taken_at = filename_date or exif_date or thm_date or media_created_date or creation_date
            stored_paths = []
            existed_at = []
            if not taken_at:
                target_path, existed = _copy_without_overwrite(
                    media_path, output_directory / "unknown date" / media_path.name
                )
                stored_paths.append(str(target_path))
                if existed:
                    existed_at.append("unknown date")
                else:
                    result.unknown_date_copies += 1
            else:
                result.dated_media += 1
                year_directory = output_directory / f"{taken_at.year:04d}"
                month_directory = year_directory / f"{taken_at.year:04d} {MONTH_NAMES[taken_at.month - 1]}"
                target_path, existed = _copy_without_overwrite(media_path, month_directory / media_path.name)
                stored_paths.append(str(target_path))
                if existed:
                    existed_at.append("month")
                else:
                    result.month_copies += 1

                album = "collage" if _is_collage(media_path, extraction_root) else _album_name(media_path, extraction_root)
                if album:
                    target_path, existed = _copy_without_overwrite(media_path, year_directory / album / media_path.name)
                    stored_paths.append(str(target_path))
                    if existed:
                        existed_at.append("album")
                    else:
                        result.album_copies += 1

            original_location = (
                f"{archive_path} :: {media_path.relative_to(extraction_root).as_posix()}"
                if archive_path else media_path.relative_to(extraction_root).as_posix()
            )
            result.report_rows.append({
                "Serial number": result.media_found,
                "File path and file name": str(media_path),
                "Found date in name": filename_date,
                "Found date in EXIF": exif_date,
                "Found date from .THM": thm_date,
                "Found date in media created": media_created_date,
                "Found date in Google Photos creation date": creation_date,
                "Where the file was at": original_location,
                "Where it has been kept now": "\n".join(stored_paths),
                "If the file already existed": (
                    f"Yes ({', '.join(existed_at)}; same filename and size)" if existed_at else "No"
                ),
            })
    return result


REPORT_COLUMNS = (
    "Serial number",
    "File path and file name",
    "Found date in name",
    "Found date in EXIF",
    "Found date from .THM",
    "Found date in media created",
    "Found date in Google Photos creation date",
    "Where the file was at",
    "Where it has been kept now",
    "If the file already existed",
)


def write_report(result: SortResult, output_directory: Path) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Media audit"
    worksheet.append(REPORT_COLUMNS)
    for row in result.report_rows:
        worksheet.append([row[column] for column in REPORT_COLUMNS])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="245A73")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column, width in zip("ABCDEFGHIJ", (14, 45, 22, 22, 22, 22, 32, 55, 70, 42)):
        worksheet.column_dimensions[column].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row[2:7]:
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
        for cell in (row[7], row[8], row[9]):
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    report_path = output_directory / "media_organization_report.xlsx"
    workbook.save(report_path)
    return report_path


def organize_takeout(
    archive_directory: Path,
    temp_directory: Path,
    output_directory: Path,
    progress_callback: Callable[[str], None] | None = None,
) -> SortResult:
    errors = []

    def log_progress(message: str) -> None:
        if message.startswith("[ERROR]"):
            errors.append(message)
        if progress_callback:
            progress_callback(message)

    archives = discover_archives(archive_directory)
    log_progress(f"Found {len(archives)} ZIP archive(s).")
    extracted_roots = extract_archives(archives, temp_directory, log_progress)
    run_directory = extracted_roots[0].parent
    result = None
    try:
        source_archives = dict(zip(extracted_roots, archives))
        result = organize_extracted(extracted_roots, output_directory, source_archives, log_progress)
        result.archives = len(archives)
        result.errors = errors
        result.report_path = write_report(result, output_directory)
        log_progress(f"Saved Excel report: {result.report_path}")
        return result
    finally:
        log_progress(f"Cleaning temporary extraction data: {run_directory}")
        try:
            shutil.rmtree(run_directory)
        except OSError as error:
            message = f"[ERROR] Could not clean temporary extraction data '{run_directory}': {error}"
            if result is not None:
                result.errors.append(message)
            log_progress(message)


class OrganizerWindow:
    def __init__(self, root: Tk) -> None:
        import tkinter as tk

        self.tk = tk
        self.root = root
        self.root.title("Google Photos Takeout Organizer")
        self.root.minsize(720, 520)
        self.events: queue.Queue[tuple[str, object]] = queue.Queue()
        self.path_vars = {
            "ZIP folder": tk.StringVar(),
            "Temporary folder": tk.StringVar(),
            "Output folder": tk.StringVar(),
        }
        self._build()
        self.root.after(150, self._poll_events)

    def _build(self) -> None:
        frame = ttk.Frame(self.root, padding=18)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(1, weight=1)
        for row, (label, variable) in enumerate(self.path_vars.items()):
            ttk.Label(frame, text=label).grid(row=row, column=0, sticky="w", padx=(0, 12), pady=8)
            ttk.Entry(frame, textvariable=variable).grid(row=row, column=1, sticky="ew", pady=8)
            ttk.Button(frame, text="Browse...", command=lambda value=variable: self._browse(value)).grid(
                row=row, column=2, padx=(10, 0), pady=8
            )
        self.status = ttk.Label(frame, text="Choose all three folders before starting.")
        self.status.grid(row=3, column=0, columnspan=3, sticky="w", pady=(12, 8))
        self.progress = ttk.Progressbar(frame, mode="indeterminate")
        self.progress.grid(row=4, column=0, columnspan=3, sticky="ew", pady=8)
        self.start_button = ttk.Button(frame, text="Organize Takeout", command=self._start)
        self.start_button.grid(row=5, column=2, sticky="e", pady=(12, 0))
        ttk.Label(frame, text="Activity log").grid(row=6, column=0, columnspan=3, sticky="w", pady=(14, 4))
        log_frame = ttk.Frame(frame)
        log_frame.grid(row=7, column=0, columnspan=3, sticky="nsew")
        frame.rowconfigure(7, weight=1)
        self.activity_log = self.tk.Text(log_frame, height=10, wrap="word", state="disabled")
        self.activity_log.tag_configure("error", foreground="#B42318")
        scrollbar = ttk.Scrollbar(log_frame, orient="vertical", command=self.activity_log.yview)
        self.activity_log.configure(yscrollcommand=scrollbar.set)
        self.activity_log.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        self.log_line_count = 0

    def _browse(self, variable: object) -> None:
        selected = filedialog.askdirectory()
        if selected:
            variable.set(selected)

    def _start(self) -> None:
        selected_paths = {label: variable.get().strip() for label, variable in self.path_vars.items()}
        if any(not value for value in selected_paths.values()):
            messagebox.showerror("Missing folder", "Select the ZIP, temporary, and output folders first.")
            return
        paths = {label: Path(value) for label, value in selected_paths.items()}
        if not paths["ZIP folder"].is_dir():
            messagebox.showerror("Invalid ZIP folder", "The selected ZIP folder does not exist.")
            return
        self.start_button.configure(state="disabled")
        self.progress.start(10)
        self.status.configure(text="Extracting archives and organizing media...")
        worker = threading.Thread(
            target=self._organize_worker,
            args=(paths["ZIP folder"], paths["Temporary folder"], paths["Output folder"]),
            daemon=True,
        )
        worker.start()

    def _organize_worker(self, archive_directory: Path, temp_directory: Path, output_directory: Path) -> None:
        try:
            result = organize_takeout(
                archive_directory,
                temp_directory,
                output_directory,
                progress_callback=lambda message: self.events.put(("log", message)),
            )
            self.events.put(("success", result))
        except Exception as error:
            self.events.put(("error", str(error)))

    def _poll_events(self) -> None:
        pending_events = []
        while True:
            try:
                pending_events.append(self.events.get_nowait())
            except queue.Empty:
                break
        if not pending_events:
            self.root.after(150, self._poll_events)
            return

        log_messages = [str(payload) for kind, payload in pending_events if kind == "log"]
        completion = next(
            ((kind, payload) for kind, payload in pending_events if kind != "log"),
            None,
        )
        if log_messages:
            self.activity_log.configure(state="normal")
            for message in log_messages:
                tag = ("error",) if message.startswith("[ERROR]") else ()
                self.activity_log.insert("end", message + "\n", *tag)
            self.log_line_count += len(log_messages)
            excess_lines = self.log_line_count - 1000
            if excess_lines > 0:
                self.activity_log.delete("1.0", f"{excess_lines + 1}.0")
                self.log_line_count -= excess_lines
            self.activity_log.see("end")
            self.activity_log.configure(state="disabled")

        if completion is None:
            if log_messages:
                self.status.configure(text=log_messages[-1])
            self.root.after(50, self._poll_events)
            return

        kind, payload = completion
        self.progress.stop()
        self.start_button.configure(state="normal")
        if kind == "success":
            result = payload
            self.status.configure(text="Organization complete.")
            messagebox.showinfo(
                "Organization complete",
                f"Archives: {result.archives}\nMedia found: {result.media_found}\n"
                f"Month copies: {result.month_copies}\nAlbum copies: {result.album_copies}\n"
                f"Unknown date copies: {result.unknown_date_copies}\n"
                f"Errors skipped: {len(result.errors)}\nExcel report: {result.report_path}",
            )
        else:
            self.status.configure(text="The operation stopped with an error.")
            messagebox.showerror("Organization failed", str(payload))
        self.root.after(150, self._poll_events)


def launch_gui() -> None:
    root = Tk()
    OrganizerWindow(root)
    root.mainloop()